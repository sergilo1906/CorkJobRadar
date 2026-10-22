"""Reference design, deterministic versions and safe preview; synthetic data only."""
import io
import zipfile
import pytest
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from test_professional_cv import candidate
from corkjobradar.autoapply.documents import render_cv, validate_document
from corkjobradar.autoapply.cv_layout import REFERENCE, cv_sections
from corkjobradar.autoapply.cv_preview import pdf_preview


@pytest.mark.parametrize('extension', ['pdf', 'docx'])
def test_repeated_render_has_same_hash_without_duplicate_drafts(candidate, extension):
    first = render_cv(candidate, extension)
    assert first == render_cv(candidate, extension)


def test_pdf_has_reference_a4_margins_and_no_clipped_text(candidate):
    import pypdfium2 as pdfium
    raw = render_cv(candidate, 'pdf')
    reader = PdfReader(io.BytesIO(raw))
    assert len(reader.pages) == 1
    page = reader.pages[0]
    assert float(page.mediabox.width) == pytest.approx(595.276, abs=.02)
    assert float(page.mediabox.height) == pytest.approx(841.89, abs=.02)
    # Inspect the actual PDF renderer's character geometry, not a story estimate.
    with pdfium.PdfDocument(raw) as pdf:
        page = pdf[0]
        text = page.get_textpage()
        try:
            boxes = [text.get_charbox(i) for i in range(text.count_chars())
                     if text.get_text_range(i, 1).strip()]
            assert boxes
            assert min(b[0] for b in boxes) >= REFERENCE.margin - 1
            assert max(b[2] for b in boxes) <= 595.276 - REFERENCE.margin + 1
            assert min(b[1] for b in boxes) >= REFERENCE.bottom - 1
            assert max(b[3] for b in boxes) <= 841.89 - REFERENCE.top + 1
        finally:
            text.close(); page.close()


def test_docx_matches_reference_colors_borders_and_typography(candidate):
    doc = Document(io.BytesIO(render_cv(candidate, 'docx')))
    page = doc.sections[0]
    assert page.page_width.mm == pytest.approx(210, abs=.01)
    assert page.left_margin.pt == REFERENCE.margin
    assert doc.styles['Title'].font.size.pt == REFERENCE.name_size
    assert str(doc.styles['Title'].font.color.rgb) == REFERENCE.navy
    assert doc.styles['Heading 1'].font.all_caps
    borders = list(doc.element.iter(qn('w:pBdr')))
    assert len(borders) >= 1 + len(cv_sections(candidate))
    assert list(doc.element.iter(qn('w:shd')))
    assert not list(doc.element.iter(qn('w:txbxContent')))
    assert not doc.tables
    for style_name in ('Normal', 'Title', 'Heading 1', 'List Bullet'):
        fonts = doc.styles[style_name].element.rPr.rFonts
        assert fonts.get(qn('w:ascii')) == 'Arial'
        assert not any(key.lower().endswith('theme') for key in fonts.attrib)
    assert doc.core_properties.author == ''
    assert doc.core_properties.last_modified_by == ''
    # PDF and DOCX share every claim, heading and contact despite formatting.
    for ext in ('pdf', 'docx'):
        normalized = ' '.join(validate_document(render_cv(candidate, ext), ext).split())
        for claim in candidate['_claims']:
            assert claim['text_en'] in normalized


def test_job_adaptation_preserves_chronology_and_promotes_only_verified_skills(candidate):
    before = [c['text_en'] for c in candidate['_claims'] if c['section'] == 'Experience']
    sections = dict(cv_sections(candidate, {'title':'Software peripherals', 'description':''}))
    assert sections['Experience'] == before
    assert 'NFQ Level 8' not in validate_document(render_cv(candidate, 'pdf'), 'pdf')


def test_preview_is_real_png_and_has_no_personal_metadata_or_file_output(candidate, tmp_path):
    from PIL import Image
    content = pdf_preview(render_cv(candidate, 'pdf'))
    assert content.startswith(b'\x89PNG\r\n\x1a\n')
    with Image.open(io.BytesIO(content)) as image:
        assert image.width > 700 and image.height > 1000
        assert not image.info
    assert not list(tmp_path.iterdir())


def test_preview_errors_are_generic_and_never_echo_document_data():
    with pytest.raises(ValueError, match='^CV_PREVIEW_UNAVAILABLE$'):
        pdf_preview(b'PRIVATE MALFORMED CONTENT')


def test_preview_rejects_annexes(candidate):
    from pypdf import PdfWriter
    reader = PdfReader(io.BytesIO(render_cv(candidate, 'pdf')))
    writer = PdfWriter(); writer.add_page(reader.pages[0]); writer.add_page(reader.pages[0])
    out = io.BytesIO(); writer.write(out)
    with pytest.raises(ValueError, match='^CV_PREVIEW_UNAVAILABLE$'):
        pdf_preview(out.getvalue())


def test_long_contacts_and_markup_remain_literal(candidate):
    candidate['full_name'] = 'SYNTHETIC Á Long Name & <b>literal</b> ' * 2
    candidate['email'] = 'synthetic.long.contact.' * 4 + '@example.invalid'
    for ext in ('pdf', 'docx'):
        text = ' '.join(validate_document(render_cv(candidate, ext), ext).split())
        assert '<b>literal</b>' in text
        assert candidate['email'] in text


def test_docx_has_no_external_relationships(candidate):
    raw = render_cv(candidate, 'docx')
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        for name in archive.namelist():
            if name.endswith('.rels'):
                assert b'TargetMode="External"' not in archive.read(name)


def test_docx_year_range_stays_whole_without_changing_the_claim(candidate):
    statement = ('SYNTHETIC: Higher Technician qualification in Multiplatform Application '
        'Development (CFGS DAM), Example Institute (2022-2024): completed; final examination passed.')
    candidate['_claims'][6]['text_en'] = statement
    doc = Document(io.BytesIO(render_cv(candidate, 'docx')))
    paragraph = next(p for p in doc.paragraphs if 'Example Institute' in p.text)
    assert '\n(2022-2024)' in paragraph.text
    assert statement in ' '.join(paragraph.text.split())


def test_docx_real_office_renderer_one_page_and_all_claims(candidate, tmp_path):
    """Optional installed LibreOffice, private owned profile, no system changes."""
    import os
    import subprocess
    import time
    from pathlib import Path
    office = Path('C:/Program Files/LibreOffice/program/soffice.com')
    if not office.is_file():
        pytest.skip('LibreOffice not installed; no claim of DOCX pagination verification')
    candidate['_claims'][6]['text_en'] = ('SYNTHETIC: Higher Technician qualification in '
        'Multiplatform Application Development (CFGS DAM), Example Institute (2022-2024): '
        'completed; final examination passed.')
    source = tmp_path/'synthetic.docx'
    source.write_bytes(render_cv(candidate, 'docx'))
    destination = tmp_path/'rendered'; destination.mkdir()
    environment = os.environ.copy()
    environment.update(TEMP=str(tmp_path), TMP=str(tmp_path))
    startup = None
    if os.name == 'nt':
        # soffice.com is a console launcher. A hidden console retains its wait/
        # pipe semantics; CREATE_NO_WINDOW can detach its underlying process.
        startup = subprocess.STARTUPINFO()
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = subprocess.SW_HIDE
    profile_argument = '-env:UserInstallation='+(tmp_path/'office-profile').as_uri()
    # Initialize a fresh, private Office profile before converting. Windows LO
    # may otherwise return from its first-run launcher without converting.
    initialized = subprocess.run([str(office), profile_argument, '--headless', '--terminate_after_init'],
        env=environment, capture_output=True, timeout=55, startupinfo=startup)
    assert initialized.returncode == 0
    result = subprocess.run([str(office), profile_argument,
        '--headless', '--norestore', '--nofirststartwizard', '--convert-to', 'pdf', '--outdir', str(destination), str(source)],
        env=environment, capture_output=True, timeout=55, startupinfo=startup)
    assert result.returncode == 0
    path = destination/'synthetic.pdf'
    deadline = time.monotonic()+40
    while not path.exists() and time.monotonic() < deadline:
        time.sleep(.1)
    assert path.is_file(), 'OFFICE_RENDER_DID_NOT_CREATE_PDF'
    reader = PdfReader(path)
    assert len(reader.pages) == 1
    text = ' '.join((reader.pages[0].extract_text() or '').split())
    assert all(c['text_en'] in text for c in candidate['_claims'])
    assert candidate['email'] in text


def test_reference_header_has_role_strapline_pdf_and_word(candidate):
    """A one-page CV must look like a CV, not just contain all text."""
    from corkjobradar.autoapply.cv_layout import _role_headline
    assert _role_headline() == 'JUNIOR IT SUPPORT | HELPDESK | SERVICE DESK'
    assert _role_headline({'title': 'QA Tester'}) == 'JUNIOR QA | SOFTWARE TESTING'
    assert _role_headline({'title': 'Junior Developer'}) == 'JUNIOR DEVELOPER | SOFTWARE DEVELOPMENT'
    assert _role_headline({'title': 'Senior Data Analyst'}) == 'JUNIOR IT SUPPORT | HELPDESK | SERVICE DESK'
    pdf = render_cv(candidate, 'pdf')
    text = validate_document(pdf, 'pdf')
    assert 'JUNIOR IT SUPPORT | HELPDESK | SERVICE DESK' in text
    assert text.index(candidate['location']) < text.index(candidate['email'])
    doc = Document(io.BytesIO(render_cv(candidate, 'docx')))
    assert doc.paragraphs[0].text == candidate['full_name']
    assert doc.paragraphs[1].text == 'JUNIOR IT SUPPORT | HELPDESK | SERVICE DESK'
    assert doc.paragraphs[1].runs[0].bold
    assert doc.paragraphs[2].text.index(candidate['location']) < doc.paragraphs[2].text.index(candidate['email'])
    assert len(PdfReader(io.BytesIO(pdf)).pages) == 1


def test_reference_skill_project_label_emphasis_is_presentation_only():
    from corkjobradar.autoapply.cv_layout import _emphasis
    for title, text in [
        ('Core Capabilities', 'IT SUPPORT: Windows, Microsoft 365, user guidance.'),
        ('Selected Projects', 'Barber Booking Platform | Astro + Firebase/Firestore'),
    ]:
        bold, tail = _emphasis(title, text, 0)
        assert bold and bold + tail == text
    bold, tail = _emphasis('Education & Training', 'Completed CFGS DAM in 2022-2024.', 0)
    assert bold + tail == 'Completed CFGS DAM in 2022-2024.'
