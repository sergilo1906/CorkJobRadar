import json
import pytest
import httpx
from corkjobradar.worker import create_run,run_worker,worker_lock,prepare_resume
from corkjobradar.models import Job
from corkjobradar.repository import save_job,reconcile_duplicate_urls
from corkjobradar.verification import verify_job,reconcile_absences
from corkjobradar.imports import import_csv,cv_skills
from corkjobradar.exporting import export_all
from corkjobradar.http import WebClient,AccessBlocked,IncompleteSource
from corkjobradar.config import now,save_profile
from openpyxl import load_workbook
from conftest import source,client

def test_absence_requires_two_full_successful_scans(db,profile):
    source(db,'greenhouse')
    initial=create_run(db)
    jid=save_job(db,'fixture',Job(external_id='1',title='IT Support',location='Cork City',url='https://fixture.test/1'),initial,profile)
    db.execute("UPDATE scan_runs SET status='completed' WHERE id=?",(initial,))
    for n in range(2):
        rid=create_run(db)
        # Esta función solo se invoca desde el trabajador tras demostrar fin completo.
        reconcile_absences(db,rid,'fixture')
        if n==0:
            assert db.rows('SELECT open_status FROM jobs')[0]['open_status']=='UNKNOWN'
        db.execute("UPDATE scan_runs SET status='completed' WHERE id=?",(rid,))
    assert db.rows('SELECT open_status FROM jobs')[0]['open_status']=='UNLISTED'

def test_verification_403_never_closed_and_form_evidence(db,profile):
    source(db,'greenhouse')
    jid=save_job(db,'fixture',Job(external_id='1',title='IT Support',url='https://fixture.test/1'),profile=profile)
    with client(profile,lambda r:httpx.Response(403)) as c:
        assert verify_job(db,jid,c)=='ACCESS_BLOCKED'
    def handle(req):
        return httpx.Response(404) if req.url.path=='/robots.txt' else httpx.Response(200,text='<form>Apply for this job <input name="resume"><button>Submit application</button></form>')
    with client(profile,handle) as c:
        assert verify_job(db,jid,c)=='OPEN_CONFIRMED'

def test_block_suspends_remaining_detail_requests(db,profile):
    source(db,'smartrecruiters',config={'company_identifier':'Fixture'})
    calls=[]
    def handle(req):
        if req.url.path=='/jobs':
            return httpx.Response(200,json={'totalFound':3,'content':[{'id':str(i),'name':'IT Support','company':{'name':'Fixture'},'location':{'city':'Cork City'},'ref':f'https://fixture.test/detail/{i}'} for i in range(3)]})
        calls.append(str(req.url)); return httpx.Response(403)
    rid=create_run(db)
    with client(profile,handle) as c: run_worker(db,rid,c)
    assert len(calls)==1
    assert len(db.rows("SELECT * FROM task_queue WHERE status='retry'"))==3
    assert db.rows('SELECT status FROM source_runs')[0]['status']=='blocked'

def test_ssrf_redirect_revalidated(profile):
    calls=[]
    def validate(url):
        if '127.0.0.1' in url: raise ValueError('Private redirect denied')
        return url
    def handle(req):
        calls.append(str(req.url)); return httpx.Response(302,headers={'location':'http://127.0.0.1/private'})
    with WebClient(profile,transport=httpx.MockTransport(handle),validator=validate,sleep=lambda _:None) as c:
        with pytest.raises(ValueError): c.get('https://fixture.test/start')
    assert len(calls)==1

def test_response_size_limit(profile):
    profile['max_response_bytes']=20
    with client(profile,lambda r:httpx.Response(200,content=b'x'*21)) as c:
        with pytest.raises(IncompleteSource): c.get('https://fixture.test/huge')

def test_robots_disallow(profile):
    def handle(req):
        return httpx.Response(200,text='User-agent: *\nDisallow: /jobs\n')
    with client(profile,handle) as c:
        with pytest.raises(AccessBlocked): c.get('https://fixture.test/jobs',check_robots=True)

def test_worker_lock_prevents_two_processes(db):
    with worker_lock(db):
        with pytest.raises(OSError):
            with worker_lock(db): pass

def test_merge_canonical_official_urls_preserves_sources_and_notes(db,profile):
    source(db,'smartrecruiters','one')
    source(db,'html','two')
    url='https://jobs.smartrecruiters.com/Fixture/123-role'
    a=save_job(db,'one',Job(external_id='1',title='Role',company='Fixture',url=url),profile=profile)
    # Simular un registro anterior a la normalización del slug.
    db.execute('UPDATE jobs SET url=? WHERE id=?',(url,a))
    b=save_job(db,'two',Job(external_id='2',title='Role',company='Fixture',url=url),profile=profile)
    assert a==b
    assert len(db.rows('SELECT * FROM job_sources'))==2
    assert reconcile_duplicate_urls(db)==[]

def test_excel_long_text_and_formula_injection(db,profile,tmp_path):
    source(db,'greenhouse')
    original='A'*40000+'END'
    save_job(db,'fixture',Job(external_id='1',title='=1+1',url='https://fixture.test/1',description=original),profile=profile)
    out=export_all(db,tmp_path)
    wb=load_workbook(out['xlsx'])
    sheet=wb['Todas_las_ofertas']
    title_column=[c.value for c in sheet[2]].index('Título')+1
    assert sheet.cell(3,title_column).value=="'=1+1"
    parts=[r[4] for r in list(wb['Textos_largos'].values)[1:] if r[0]=='Todas_las_ofertas' and r[2]=='description']
    assert ''.join(parts)==original
    wb.close()

def test_csv_real_import_states_and_profile_validation(db,profile):
    ids=import_csv(db,'title,url,company,location,description,detail_complete\nIT Support,https://fixture.test/1,Fixture,Cork City,Brief,False\n',profile)
    assert len(ids)==1 and db.rows('SELECT analysis_status FROM jobs')[0]['analysis_status']=='PARCIAL_POR_DATOS'
    broken=dict(profile); broken['weights']={**profile['weights'],'salary':50}
    with pytest.raises(ValueError): save_profile(broken)

def test_cv_docx_local_only():
    import io
    from docx import Document
    stream=io.BytesIO(); doc=Document(); doc.add_paragraph('SYNTHETIC CV: Java, SQL, warehouse.'); doc.save(stream)
    assert set(cv_skills(stream.getvalue(),'synthetic.docx'))=={'Java','SQL','warehouse'}
