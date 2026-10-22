import json
from pathlib import Path
import httpx
from openpyxl import load_workbook
from corkjobradar.worker import create_run,run_worker,prepare_resume
from corkjobradar.audit import coverage
from corkjobradar.exporting import export_all
from corkjobradar.db import Database
from corkjobradar.config import now
from corkjobradar.imports import import_url
from conftest import source,client,fixture

def synthetic_text(index):
    text = ('SYNTHETIC TEST ONLY. Provide support for colleagues, troubleshoot computers and maintain records. '
            'Install equipment, assist staff with accounts, document incidents and collaborate with a team. '
            'Training provided. Monday to Friday. Permanent position. At least 1 year experience. '
            'No vacancy in this test is real. ')
    if index>=8:
        text += 'Salary EUR 32000 per year. '
    if index>=134:
        text += 'This job has been closed. '
    return text

def setup_large(db):
    source(db,'smartrecruiters','long','https://fixture.test/smart',{'page_size':20,'company_identifier':'Fixture'})
    source(db,'greenhouse','duplicates','https://fixture.test/greenhouse')

def large_handler(failed=True,interrupt_at=None,seen=None):
    def handle(req):
        path=req.url.path
        if seen is not None:
            seen.append(str(req.url))
        if path=='/smart':
            offset=int(req.url.params.get('offset',0))
            if interrupt_at is not None and offset>=interrupt_at:
                raise KeyboardInterrupt('SYNTHETIC INTERRUPTION')
            rows=[{'id':str(i),'name':f'IT Support {i}','company':{'name':'SYNTHETIC TEST ONLY'},
                   'location':{'city':'Cork City','country':'ie'},'ref':f'https://fixture.test/detail/{i}',
                   'postingUrl':f'https://fixture.test/official/{i}'} for i in range(offset,min(offset+20,137))]
            return httpx.Response(200,json={'totalFound':137,'content':rows})
        if path.startswith('/detail/'):
            index=int(path.split('/')[-1])
            if failed and 120<=index<125:
                return httpx.Response(503)
            return httpx.Response(200,json={'id':str(index),'name':f'IT Support {index}',
                'company':{'name':'SYNTHETIC TEST ONLY'},'location':{'city':'Cork City','country':'ie'},
                'postingUrl':f'https://fixture.test/official/{index}',
                'jobAd':{'sections':{'jobDescription':{'title':'Duties','text':synthetic_text(index)}}}})
        if path=='/greenhouse':
            rows=[{'id':f'dup-{i}','title':f'IT Support {i}','company_name':'SYNTHETIC TEST ONLY',
                   'location':{'name':'Cork City, ie'},'absolute_url':f'https://fixture.test/official/{i}',
                   'content':synthetic_text(i)} for i in range(12)]
            return httpx.Response(200,json={'jobs':rows,'meta':{'total':12}})
        raise AssertionError(f'Unexpected URL: {req.url}')
    return handle

def test_137_offers_7_pages_12_duplicates_5_failures(db,profile,tmp_path):
    setup_large(db)
    run_id=create_run(db)
    with client(profile,large_handler(),db) as c:
        run_worker(db,run_id,c)
    report=coverage(db,run_id)
    assert report['observed_postings']==149
    assert report['unique_jobs']==137
    assert report['duplicate_publications']==12
    assert report['target_or_pending_jobs']==137
    assert report['full_descriptions_in_scope']==132
    assert report['analysis_groups']=={'basico_o_parcial':132,'pendiente':5}
    assert report['reconciled']
    assert len(db.rows("SELECT * FROM jobs WHERE open_status='CLOSED_CONFIRMED'"))==3
    assert len(db.rows("SELECT * FROM task_queue WHERE status='retry'"))==5
    assert db.rows("SELECT pages_visited FROM source_runs WHERE source_id='long'")[0]['pages_visited']==7
    assert db.rows('SELECT status FROM scan_runs')[0]['status']=='partial'
    # Candidatura y exportación reales sobre fixtures inequívocos fuera de producción.
    job_id=db.rows('SELECT id FROM jobs LIMIT 1')[0]['id']
    db.execute('INSERT INTO user_actions(job_id,action,notes,application_date,created_at) VALUES(?,?,?,?,?)',
               (job_id,'solicitada','SYNTHETIC TEST ONLY','2026-10-19',now()))
    out=export_all(db,tmp_path/'exports',run_id)
    wb=load_workbook(out['xlsx'],read_only=True)
    assert wb['Todas_las_ofertas'].max_row==139
    assert wb['Candidaturas'].max_row==3
    assert {'Fuera_de_alcance','Errores','Auditoria','Historial','Accesibilidad_pendiente'}<=set(wb.sheetnames)
    wb.close()
    assert Path(out['csv']).read_bytes().startswith(b'\xef\xbb\xbf')
    assert len(json.loads(Path(out['json']).read_text(encoding='utf-8'))['tables']['Todas_las_ofertas'])==137
    # Reinicio: únicamente se descargan los cinco detalles pendientes.
    restarted=Database(db.path)
    prepare_resume(restarted,run_id)
    requests=[]
    with client(profile,large_handler(failed=False,seen=requests),db) as c:
        run_worker(restarted,run_id,c)
    assert len(requests)==5
    assert len(restarted.rows('SELECT * FROM jobs'))==137
    assert restarted.rows('SELECT status FROM scan_runs')[0]['status']=='completed'
    assert coverage(restarted,run_id)['analysis_groups']=={'basico_o_parcial':137}
    assert restarted.rows('SELECT action FROM user_actions')[0]['action']=='solicitada'

def test_interrupt_resume_from_persistent_cursor(db,profile):
    setup_large(db)
    # Recorrer solo la fuente larga para poder observar el cursor al interrumpir.
    run_id=create_run(db,['long'])
    try:
        with client(profile,large_handler(failed=False,interrupt_at=60),db) as c:
            run_worker(db,run_id,c)
    except KeyboardInterrupt:
        pass
    assert len(db.rows('SELECT * FROM jobs'))==60
    assert db.rows('SELECT cursor FROM source_runs')[0]['cursor']=='60'
    prepare_resume(db,run_id)
    requests=[]
    with client(profile,large_handler(failed=False,seen=requests),db) as c:
        run_worker(db,run_id,c)
    assert len(db.rows('SELECT * FROM jobs'))==137
    assert not any('/smart?' in url and 'offset=0&' in url for url in requests)
    assert coverage(db,run_id)['reconciled']

def test_import_omission_persisted_on_403(db,profile):
    with client(profile,lambda req:httpx.Response(403),db) as c:
        ids=import_url(db,'https://fixture.test/omitted',c,profile)
    assert len(ids)==1
    assert db.rows('SELECT * FROM audit_events WHERE kind=?',('INVESTIGAR_OMISION',))
    assert db.rows('SELECT open_status FROM jobs')[0]['open_status']=='ACCESS_BLOCKED'
    assert len(Database(db.path).rows('SELECT * FROM jobs'))==1

def test_import_url_jsonld(db,profile):
    def handle(req):
        return httpx.Response(404) if req.url.path=='/robots.txt' else httpx.Response(200,text=fixture('jobposting.html'))
    with client(profile,handle,db) as c:
        ids=import_url(db,'https://fixture.test/job/html-test',c,profile)
    assert len(ids)==1
    assert len(db.rows('SELECT * FROM job_analyses'))==1
