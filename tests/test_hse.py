import io
import json
import httpx
import pytest
from docx import Document
from conftest import client, fixture, source
from corkjobradar.adapters import HSE
from corkjobradar.documents import document_text

BASE = 'https://about.hse.ie/jobs/job-search/?county=Cork&page=1'


def test_hse_pagination_from_captured_public_pages(db, profile):
    def handler(request):
        if request.url.path == '/robots.txt':
            return httpx.Response(200, text='User-agent: *\nDisallow: /api/v2/download-file/')
        name = 'hse-page2.html' if request.url.params.get('page') == '2' else 'hse-page1.html'
        return httpx.Response(200, text=fixture(name))
    adapter = HSE(source(db, 'hse', url=BASE), client(profile, handler), profile)
    pages = list(adapter.discover())
    assert [len(p.jobs) for p in pages] == [10, 5]
    assert pages[-1].expected_total == 15 and pages[-1].complete is True
    assert len({j.url for p in pages for j in p.jobs}) == 15
    assert all(not j.detail_complete for p in pages for j in p.jobs)
    assert len(list(adapter.discover(pages[0].cursor))) == 1


def test_hse_changed_markup_cannot_claim_complete(db, profile):
    def handler(request):
        return httpx.Response(200, text='User-agent: *' if request.url.path == '/robots.txt' else '<main>No cards</main>')
    with pytest.raises(RuntimeError, match='total/rango'):
        list(HSE(source(db, 'hse', url=BASE), client(profile, handler), profile).discover())


def test_hse_preserves_unreadable_attachment_as_pending(db, profile):
    def handler(request):
        if request.url.path == '/robots.txt':
            return httpx.Response(200, text='User-agent: *')
        if request.url.path.endswith('.docx'):
            return httpx.Response(403)
        return httpx.Response(200, text=fixture('hse-detail.html'))
    adapter = HSE(source(db, 'hse', url=BASE), client(profile, handler), profile)
    job = adapter.fetch_detail(dict(external_id='NRS15623', title='Technical Specialist', url=BASE))
    assert not job.detail_complete and len(job.description) > 1000
    assert sum(a['status']=='REVISION_MANUAL' for a in job.raw['attachments']) == 2
    assert 'Gurranabraher, Cork' in job.location
    assert job.raw['application_deadline']=='2026-11-05T15:00:00'
    assert job.raw['deadline_timezone']=='Europe/Dublin'
    assert job.raw['application_deadline_original']=='05/11/2026  15:00:00'


def test_docx_reads_tables_and_rejects_unknown_format():
    doc = Document()
    doc.add_paragraph('SYNTHETIC UNIT TEST ONLY. ' + 'description ' * 35)
    doc.add_table(rows=1, cols=1).cell(0,0).text = 'Salary EUR 15 per hour'
    out = io.BytesIO(); doc.save(out)
    assert 'Salary EUR 15 per hour' in document_text(out.getvalue(), 'spec.docx')
    with pytest.raises(ValueError, match='Formato'):
        document_text(b'legacy doc binary', 'spec.doc')
