import io
import json
import pytest
from openpyxl import load_workbook
from corkjobradar.analysis import analyze_job
from corkjobradar.autoapply.academic import assessment,review_requirement,dam_state
from corkjobradar.autoapply.manual_requirements import checklist
from corkjobradar.autoapply.vault import approve_fact,candidate_snapshot
from corkjobradar.autoapply.selection_queue import refresh_queue,prepare_all_local
from corkjobradar.reevaluation import reevaluate_all,opportunity_report,report_excel,candidate_snapshot_hash
from scripts.update_dam_completed_20261022 import FIELD,COMPLETED
from test_usable_it_workspace import case


@pytest.mark.parametrize('quote,code,status',[
    ('CFGS DAM required.','DAM_COMPLETION_CONFIRMED','CUMPLIDO'),
    ('Completed IT qualification required.','DAM_COMPLETION_CONFIRMED','CUMPLIDO'),
    ('Bachelor degree required.','ACADEMIC_EQUIVALENCE_PENDING','PENDIENTE'),
    ('Honours Degree at NFQ Level 8 required.','ACADEMIC_EQUIVALENCE_PENDING','PENDIENTE'),
    ('Honours degree or Honours Bachelor required.','ACADEMIC_EQUIVALENCE_PENDING','PENDIENTE'),
    ('NFQ Level 6 qualification required.','ACADEMIC_EQUIVALENCE_PENDING','PENDIENTE'),
    ('Degree or equivalent experience required.','ACADEMIC_ALTERNATIVE_PENDING','PENDIENTE'),
    ('Currently pursuing a Bachelor degree.','CURRENT_STUDENT_STATUS_PENDING','PENDIENTE'),
    ('Currently enrolled in a software degree required.','CURRENT_STUDENT_STATUS_PENDING','PENDIENTE'),
    ('Diploma in network engineering required.','RELATED_QUALIFICATION_REVIEW','PENDIENTE'),
    ('CFGS DAM and fluent German required.','RELATED_QUALIFICATION_REVIEW','PENDIENTE'),
    ('CFGS DAM with distinction required.','RELATED_QUALIFICATION_REVIEW','PENDIENTE'),
])
def test_approved_qualification_never_silently_accredits_equivalence_or_extra_conditions(db,quote,code,status):
    approve_fact(db,FIELD,COMPLETED,'SYNTHETIC qualification only',True)
    result=review_requirement(quote,candidate_snapshot(db))
    assert result['code']==code and result['status']==status and not result['equivalence_assumed']
    assert bool(result['fact_ids'])==(status=='CUMPLIDO')


def test_previous_dam_incomplete_never_passes_and_preferred_remains_optional(db,profile):
    job=case(db,profile,'Completed IT qualification required.\nBachelor degree preferred.')
    approve_fact(db,FIELD,{'section':'Education','text_en':'CFGS DAM: ongoing; MP13 remaining.'},'SYNTHETIC',True)
    before=analyze_job(db,job['id'],profile);rows=checklist(db,job)
    assert rows[0]['status']=='PENDIENTE' and rows[1]['level']=='DESEABLE'
    previous=db.rows('SELECT * FROM job_analyses')
    approve_fact(db,FIELD,COMPLETED,'SYNTHETIC changed qualification',True)
    after=analyze_job(db,job['id'],profile)
    assert after['academic_review']['dam']['state']=='COMPLETED'
    assert after['candidate_facts_hash']!=before['candidate_facts_hash']
    assert checklist(db,job)[0]['status']=='CUMPLIDO'
    assert checklist(db,job)[1]['status']=='PENDIENTE'
    for r in previous:assert db.rows('SELECT * FROM job_analyses WHERE id=?',(r['id'],))[0]==r
    assert len(db.rows('SELECT * FROM job_analyses'))==2
    analyze_job(db,job['id'],profile)
    assert len(db.rows('SELECT * FROM job_analyses'))==2


def test_all_records_including_outside_and_non_it_are_reevaluated_and_resume_without_duplication(db,profile,monkeypatch):
    from corkjobradar.repository import save_job
    from corkjobradar.models import Job
    job=case(db,profile,'3 years of IT support experience required.\nDegree in computing required.')
    from conftest import source
    source(db,'html',sid='academic_fixture')
    approve_fact(db,FIELD,COMPLETED,'SYNTHETIC',True)
    save_job(db,'academic_fixture',Job(external_id='other',title='Warehouse Operative',company='SYNTHETIC OTHER',location='Dublin',url='https://fixture.test/other',description='SYNTHETIC stock handling. '*50,detail_complete=True),profile=profile)
    original=db.rows('SELECT * FROM candidate_documents');calls=[]
    import httpx
    monkeypatch.setattr(httpx.Client,'request',lambda *a,**k:(_ for _ in ()).throw(AssertionError('NO_HTTP_REANALYSIS')))
    stopped=reevaluate_all(db,should_stop=lambda:len(db.rows('SELECT * FROM job_analyses'))==1)
    assert stopped['paused'] and stopped['jobs_processed']==1
    done=reevaluate_all(db)
    assert done['jobs_processed']==done['jobs_seen']==2 and done['all_stored_jobs_processed']
    assert len(db.rows('SELECT * FROM job_analyses'))==2
    again=reevaluate_all(db);assert again['new_analyses']==0
    assert candidate_snapshot_hash(db)==candidate_snapshot(db)['_cv_facts_hash']
    refresh_queue(db);report=opportunity_report(db)
    assert report['all_jobs']==2 and report['it_jobs']==1 and report['reasonable_local_opportunities']==1
    assert not report['cv_approved'] and report['current_academic_analyses']==2
    assert report['offers'][0]['academic_pending']
    assert db.rows('SELECT * FROM candidate_documents')==original
    prepared=prepare_all_local(db)
    assert prepared['manual_ready']==prepared['submitted']==0
    assert not db.rows('SELECT * FROM manual_application_handoffs')
    assert not db.rows('SELECT * FROM application_attempts')
    book=load_workbook(io.BytesIO(report_excel(report)))
    assert book['Oportunidades IT'].max_row==2 and book['Alcance']['B5'].value==report['limits']


def test_no_reanalysis_worker_can_run_while_search_has_reserved_work(db):
    from corkjobradar.worker import create_run
    create_run(db)
    with pytest.raises(RuntimeError,match='SEARCH_ALREADY_ACTIVE'):reevaluate_all(db)


def test_completed_dam_keyword_must_be_in_education_not_web_prompt_injection():
    candidate={'_claims':[{'id':1,'section':'Skills','text_en':'CFGS DAM completed'}]}
    assert dam_state(candidate)['state']=='UNKNOWN'
    candidate['_claims'][0]['section']='Education';candidate['_claims'][0]['text_en']='CFGS DAM not completed'
    assert dam_state(candidate)['state']=='ONGOING'


def test_scheduled_work_reanalyses_new_arrivals_and_waits_for_document_approval(db,profile,monkeypatch):
    from corkjobradar import runtime as rt
    from corkjobradar.repository import save_job
    from corkjobradar.models import Job
    from conftest import source
    case(db,profile,'Completed IT qualification required.')
    approve_fact(db,FIELD,COMPLETED,'SYNTHETIC',True)
    rt.update(db,desired='RUNNING',state='PREPARING',auto_prepare=1)
    monkeypatch.setattr('corkjobradar.autoapply.selection_queue.prepare_all_local',lambda *a,**k:(_ for _ in ()).throw(AssertionError('CV_MUST_NOT_BE_APPROVED_BY_WORKER')))
    rt.run_work(db,'PREPARE')
    assert len(db.rows('SELECT * FROM job_analyses'))==1
    source(db,'html',sid='new_source')
    save_job(db,'new_source',Job(external_id='new',title='Application Support',company='SYNTHETIC NEW',location='Cork',url='https://fixture.test/new',description='Provide application support and troubleshoot software incidents. '*20,detail_complete=True),profile=profile)
    rt.run_work(db,'PREPARE')
    assert len(db.rows('SELECT * FROM job_analyses'))==2
    assert len(db.rows('SELECT * FROM it_application_queue'))==2
    assert rt.state(db)['interval_hours']==6 and rt.state(db)['desired']=='RUNNING'
    assert not db.rows('SELECT * FROM candidate_documents') and not db.rows('SELECT * FROM manual_application_handoffs')
    assert not db.rows('SELECT * FROM application_attempts')
