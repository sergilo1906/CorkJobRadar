import hashlib
import io
from pathlib import Path
import zipfile
import re
from xml.sax.saxutils import escape
from ..config import now,digest
from .vault import protect,private_folder,candidate_snapshot
from .policy import event


def validate_document(content, extension):
    try:return _validate_document(content,extension)
    except Exception:
        # Parser exceptions can contain source text or filenames. Keep them out of UI/logs.
        raise ValueError('DOCUMENT_VALIDATION_FAILED') from None


def _validate_document(content, extension):
    if len(content)>10_000_000:raise ValueError('DOCUMENT_TOO_LARGE')
    if extension=='pdf':
        from pypdf import PdfReader
        if not content.startswith(b'%PDF-'):raise ValueError('INVALID_PDF_HEADER')
        reader=PdfReader(io.BytesIO(content))
        if reader.is_encrypted or not 1<=len(reader.pages)<=20:raise ValueError('INVALID_PDF_PAGES')
        for page in reader.pages:
            if page.get('/AA') or page.get('/Annots'):raise ValueError('ACTIVE_PDF_CONTENT_REJECTED')
        if any(k in reader.trailer['/Root'] for k in ('/OpenAction','/AA','/AcroForm','/Names')):raise ValueError('ACTIVE_PDF_CONTENT_REJECTED')
        text='\n'.join(p.extract_text() or '' for p in reader.pages)
    elif extension=='docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist())>30_000_000:raise ValueError('EXPANDED_DOCX_TOO_LARGE')
            if any('vbaProject' in n or '/embeddings/' in n for n in z.namelist()):raise ValueError('ACTIVE_DOCX_CONTENT_REJECTED')
            for n in z.namelist():
                if n.endswith('.rels'):
                    from lxml import etree
                    tree=etree.fromstring(z.read(n),etree.XMLParser(resolve_entities=False,no_network=True))
                    if any(e.get('TargetMode','').lower()=='external' for e in tree.iter()):raise ValueError('EXTERNAL_DOCX_CONTENT_REJECTED')
        text='\n'.join(p.text for p in Document(io.BytesIO(content)).paragraphs)
    else:raise ValueError('PDF_OR_DOCX_REQUIRED')
    if not text.strip():raise ValueError('EMPTY_DOCUMENT')
    return text


def render_cv(candidate, extension='docx', job=None):
    if job:
        from .it_scope import require_it_job
        require_it_job(job)
    if not candidate.get('full_name') or not candidate.get('email') or not candidate.get('_claims'):
        raise ValueError('VERIFIED_CONTACT_AND_CLAIMS_REQUIRED')
    # Fixed English labels + approved English claims; no web text and no LLM instructions.
    lines=[candidate['full_name'],candidate['email']]
    for key in ('phone','location','linkedin_url'):
        if candidate.get(key):lines.append(candidate[key])
    for section in ('Experience','Education','Skills','Languages','Availability'):
        claims=[c['text_en'] for c in candidate['_claims'] if c['section']==section]
        if job:
            # Reorder approved claims by literal job vocabulary. Web instructions never become CV content.
            vocabulary=set(re.findall(r'[a-z]{3,}',(job.get('title','')+' '+job.get('description','')).lower()))
            claims.sort(key=lambda c:-len(vocabulary.intersection(re.findall(r'[a-z]{3,}',c.lower()))))
        if claims:lines.extend([section,*claims])
    return _render_lines(lines,candidate,extension)


def render_letter(candidate,job,extension='docx'):
    from .it_scope import require_it_job
    require_it_job(job)
    if not candidate.get('full_name') or not candidate.get('email') or not candidate.get('_claims'):
        raise ValueError('VERIFIED_CONTACT_AND_CLAIMS_REQUIRED')
    lines=[candidate['full_name'],candidate['email'],'Dear Hiring Team,',
           'Please consider my application for '+job['title']+' at '+job['company']+'.',
           'My verified background includes:',*[c['text_en'] for c in candidate['_claims']],
           'I would welcome the opportunity to discuss this position.','Kind regards,',candidate['full_name']]
    return _render_lines(lines,candidate,extension)


def _render_lines(lines,candidate,extension):
    if extension=='docx':
        from docx import Document
        doc=Document()
        for n,line in enumerate(lines):doc.add_paragraph(line,style='Title' if n==0 else 'Normal')
        buffer=io.BytesIO();doc.save(buffer);content=buffer.getvalue()
    elif extension=='pdf':
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        font=Path('C:/Windows/Fonts/arial.ttf')
        if not font.is_file():raise RuntimeError('UNICODE_FONT_UNAVAILABLE')
        if 'RadarArial' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('RadarArial',str(font)))
        buffer=io.BytesIO();styles=getSampleStyleSheet();styles['Normal'].fontName='RadarArial'
        story=[]
        for line in lines:story.extend([Paragraph(escape(line).replace('\n','<br/>'),styles['Normal']),Spacer(1,8)])
        SimpleDocTemplate(buffer).build(story);content=buffer.getvalue()
    else:raise ValueError('PDF_OR_DOCX_REQUIRED')
    text=validate_document(content,extension)
    for claim in candidate['_claims']:
        if ' '.join(claim['text_en'].split()) not in ' '.join(text.split()):raise ValueError('GENERATED_DOCUMENT_LOST_A_CLAIM')
    return content


def store_document(db, content, extension, approved=False, kind='CV', claim_ids=None):
    text=validate_document(content,extension)
    if kind=='CV' and extension=='pdf':
        from pypdf import PdfReader
        if len(PdfReader(io.BytesIO(content)).pages)!=1:
            raise ValueError('CV_FIRST_PAGE_ONLY_SEPARATE_ANNEXES')
    facts=candidate_snapshot(db)
    ids=claim_ids if claim_ids is not None else [c['id'] for c in facts['_claims']]
    all_claims={c['id']:c for c in facts['_claims']}
    if not ids or not set(ids).issubset(all_claims):raise ValueError('DOCUMENT_REQUIRES_APPROVED_CLAIMS')
    # Imported files need explicit human review of all text; no claim of automatic semantic verification.
    if not approved:raise ValueError('EXPLICIT_DOCUMENT_REVIEW_REQUIRED')
    for fid in ids:
        if ' '.join(all_claims[fid]['text_en'].split()) not in ' '.join(text.split()):raise ValueError('DOCUMENT_CLAIM_NOT_PRESENT')
    file_hash=hashlib.sha256(content).hexdigest();folder=private_folder(db)
    path=folder/(file_hash+'.'+extension+'.dpapi')
    if not path.exists():path.write_bytes(protect(content))
    doc_id=db.execute('''INSERT INTO candidate_documents(kind,encrypted_path,file_hash,facts_hash,claim_ids,created_at,approved_at)
        VALUES(?,?,?,?,?,?,?)''',(kind,str(path),file_hash,facts['_cv_facts_hash'],__import__('json').dumps(ids),now(),now()))
    event(db,'DOCUMENT_APPROVED',{'document_id':doc_id,'hash':file_hash})
    return doc_id


def document_bytes(db, document_id):
    row=db.rows('SELECT * FROM candidate_documents WHERE id=? AND revoked_at IS NULL AND approved_at IS NOT NULL',(document_id,))
    if not row:raise ValueError('APPROVED_DOCUMENT_MISSING')
    row=row[0]
    if row['facts_hash']!=candidate_snapshot(db)['_cv_facts_hash']:raise ValueError('DOCUMENT_FACTS_CHANGED')
    path=Path(row['encrypted_path']).resolve()
    if not path.is_relative_to(private_folder(db)):raise ValueError('DOCUMENT_PATH_OUTSIDE_VAULT')
    content=protect(path.read_bytes(),decrypt=True)
    if hashlib.sha256(content).hexdigest()!=row['file_hash']:raise ValueError('DOCUMENT_HASH_MISMATCH')
    extension=path.name.split('.')[-2]
    validate_document(content,extension)
    return content,extension,row['file_hash']


def create_cv_review(db, content, extension):
    """Persist an encrypted draft. Generating/importing never records document consent."""
    text=validate_document(content,extension)
    if extension=='pdf':
        from pypdf import PdfReader
        if len(PdfReader(io.BytesIO(content)).pages)!=1:raise ValueError('CV_FIRST_PAGE_ONLY_SEPARATE_ANNEXES')
    facts=candidate_snapshot(db)
    if not facts.get('full_name') or not facts.get('email') or not facts['_claims']:
        raise ValueError('VERIFIED_CONTACT_AND_CLAIMS_REQUIRED')
    for claim in facts['_claims']:
        if ' '.join(claim['text_en'].split()) not in ' '.join(text.split()):raise ValueError('DOCUMENT_CLAIM_NOT_PRESENT')
    sha=hashlib.sha256(content).hexdigest()
    with db.transaction():
        existing=db.rows("SELECT id FROM candidate_documents WHERE kind='CV' AND file_hash=? AND facts_hash=? AND revoked_at IS NULL ORDER BY id DESC LIMIT 1",(sha,facts['_cv_facts_hash']))
        if existing:return existing[0]['id']
        path=private_folder(db)/(sha+'.'+extension+'.dpapi')
        if not path.exists():path.write_bytes(protect(content))
        did=db.execute('''INSERT INTO candidate_documents(kind,encrypted_path,file_hash,facts_hash,claim_ids,created_at)
                          VALUES('CV',?,?,?,?,?)''',(str(path),sha,facts['_cv_facts_hash'],__import__('json').dumps([c['id'] for c in facts['_claims']]),now()))
        event(db,'CV_REVIEW_CREATED',{'document_id':did,'hash':sha})
        return did


def cv_review_bytes(db, document_id):
    rows=db.rows("SELECT * FROM candidate_documents WHERE id=? AND kind='CV' AND revoked_at IS NULL",(document_id,))
    if not rows:raise ValueError('CV_REVIEW_NOT_FOUND')
    row=rows[0]
    facts=candidate_snapshot(db)
    if row['facts_hash']!=facts['_cv_facts_hash']:raise ValueError('DOCUMENT_FACTS_CHANGED')
    path=Path(row['encrypted_path']).resolve()
    if not path.is_relative_to(private_folder(db)):raise ValueError('DOCUMENT_PATH_OUTSIDE_VAULT')
    content=protect(path.read_bytes(),decrypt=True)
    if hashlib.sha256(content).hexdigest()!=row['file_hash']:raise ValueError('DOCUMENT_HASH_MISMATCH')
    extension=path.name.split('.')[-2]
    text=validate_document(content,extension)
    import json
    if set(json.loads(row['claim_ids']))!={c['id'] for c in facts['_claims']}:raise ValueError('DOCUMENT_CLAIMS_CHANGED')
    if any(' '.join(c['text_en'].split()) not in ' '.join(text.split()) for c in facts['_claims']):raise ValueError('DOCUMENT_CLAIM_NOT_PRESENT')
    return content,extension,row['file_hash'],text


def approve_cv_review(db, document_id, expected_hash, expected_facts_hash, confirmed=False):
    if not confirmed:raise ValueError('EXPLICIT_DOCUMENT_REVIEW_REQUIRED')
    with db.transaction() as con:
        if not con.in_transaction:con.execute('BEGIN IMMEDIATE')
        _,_,sha,_=cv_review_bytes(db,document_id)
        if expected_hash!=sha or expected_facts_hash!=candidate_snapshot(db)['_facts_hash']:raise ValueError('DOCUMENT_OR_FACTS_CHANGED')
        row=db.rows('SELECT approved_at FROM candidate_documents WHERE id=?',(document_id,))[0]
        if row['approved_at']:return document_id
        db.execute('UPDATE candidate_documents SET approved_at=? WHERE id=?',(now(),document_id))
        from .policy import pause
        if db.rows('SELECT enabled FROM autoapply_control')[0]['enabled']:pause(db,'CV_VERSION_CHANGED_REAUTHORIZE_POLICY',alert=True)
        event(db,'CV_REVIEW_APPROVED',{'document_id':document_id,'hash':sha})
        return document_id
