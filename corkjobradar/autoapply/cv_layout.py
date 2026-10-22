"""Single CV presentation model; never creates facts, consent or database writes."""
import io
import re
from pathlib import Path
from xml.sax.saxutils import escape


SECTIONS = (
    ('Profile', 'Professional Profile'),
    ('Skills', 'Core Capabilities'),
    ('Experience', 'Experience'),
    ('Projects', 'Selected Projects'),
    ('Education', 'Education & Training'),
    ('Additional Experience', 'Additional Experience'),
    ('Languages', 'Languages'),
    ('Availability', 'Availability'),
)


def cv_sections(candidate, job=None):
    """Keep every approved claim verbatim; job text can affect ordering only."""
    groups = {key: [] for key, _ in SECTIONS}
    vocabulary = set(re.findall(r'[a-z]{3,}', (
        (job or {}).get('title', '') + ' ' + (job or {}).get('description', '')).lower()))
    for claim in candidate['_claims']:
        section = claim['section']
        if section not in groups:
            raise ValueError('UNSUPPORTED_CV_SECTION')
        text = claim['text_en']
        # Older approved profiles stored project descriptions among skills.
        # Move complete claims only; do not paraphrase or infer missing experience.
        if section == 'Skills' and re.search(
                r'\b(?:project|booking platform|data automation)\b', text, re.I):
            section = 'Projects'
        groups[section].append(text)
    result = []
    for key, label in SECTIONS:
        claims = groups[key]
        if vocabulary:
            claims = sorted(claims, key=lambda text: -len(vocabulary.intersection(
                re.findall(r'[a-z]{3,}', text.lower()))))
        if claims:
            result.append((label, claims))
    return result


def _font_files():
    pairs = (
        ('C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/arialbd.ttf'),
        ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
         '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
    )
    for regular, bold in pairs:
        if Path(regular).is_file() and Path(bold).is_file():
            return regular, bold
    raise RuntimeError('UNICODE_FONT_UNAVAILABLE')


def render_professional_cv(candidate, extension, job=None):
    sections = cv_sections(candidate, job)
    contact = ' | '.join(candidate[k] for k in ('email', 'phone', 'location') if candidate.get(k))
    linkedin = candidate.get('linkedin_url')
    if extension == 'pdf':
        return _pdf(candidate['full_name'], contact, linkedin, sections)
    if extension == 'docx':
        return _docx(candidate['full_name'], contact, linkedin, sections)
    raise ValueError('PDF_OR_DOCX_REQUIRED')


def _pdf(name, contact, linkedin, sections):
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from pypdf import PdfReader
    regular, bold = _font_files()
    for font_name, path in (('RadarCV', regular), ('RadarCVBold', bold)):
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_name, path))
    # Try readable sizes only. Overflow is explicit; never truncate approved facts.
    for size in (10, 9.5, 9):
        body = ParagraphStyle('CVBody', fontName='RadarCV', fontSize=size,
                              leading=size*1.25, spaceAfter=3)
        heading = ParagraphStyle('CVSection', parent=body, fontName='RadarCVBold',
                                 fontSize=10, spaceBefore=8, spaceAfter=4, keepWithNext=True)
        title = ParagraphStyle('CVName', parent=body, fontName='RadarCVBold',
                               fontSize=22, leading=26, spaceAfter=5)
        small = ParagraphStyle('CVContact', parent=body, fontSize=9, leading=12)
        paragraph = lambda text, style: Paragraph(escape(text).replace('\n', '<br/>'), style)
        story = [paragraph(name, title), paragraph(contact, small)]
        if linkedin:
            story.append(paragraph(linkedin, small))
        story.append(Spacer(1, 3))
        for label, claims in sections:
            story.append(paragraph(label, heading))
            story.extend(paragraph(text, body) for text in claims)
        buffer = io.BytesIO()
        SimpleDocTemplate(buffer, pagesize=A4, leftMargin=40, rightMargin=40,
                          topMargin=32, bottomMargin=32, title='Curriculum vitae',
                          author=name).build(story)
        content = buffer.getvalue()
        if len(PdfReader(io.BytesIO(content)).pages) == 1:
            return content
    raise ValueError('CV_CONTENT_EXCEEDS_ONE_PAGE_REVIEW_REQUIRED')


def _docx(name, contact, linkedin, sections):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.oxml.ns import qn
    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Mm(210), Mm(297)
    page.left_margin = page.right_margin = Mm(14)
    page.top_margin = page.bottom_margin = Mm(11.3)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size in (('Title', 22), ('Heading 1', 10)):
        style = doc.styles[style_name]
        style.font.name, style.font.size = 'Arial', Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_after = Pt(4)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    # Word's built-in Title style can inherit a decorative blue paragraph rule.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.add_paragraph(name, style='Title')
    for text in (contact, linkedin):
        if text:
            p = doc.add_paragraph(text)
            for run in p.runs:
                run.font.size = Pt(9)
    for label, claims in sections:
        doc.add_paragraph(label, style='Heading 1')
        for text in claims:
            doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
