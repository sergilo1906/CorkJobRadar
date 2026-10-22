"""Append-only local reanalysis of ALL stored jobs; restart reuses completed versions."""
from collections import Counter
import json
from .config import load_profile,now
from .analysis import analyze_job
from .autoapply.vault import candidate_snapshot
from .autoapply.academic import assessment,dam_state


def reevaluate_all(db,*,should_stop=None):
    from .worker import worker_lock
    with worker_lock(db):
        if db.rows("SELECT id FROM scan_runs WHERE status='running'"):raise RuntimeError('SEARCH_ALREADY_ACTIVE')
        candidate=candidate_snapshot(db);profile=load_profile();profile['analyzer']='basic'
        ids=[r['id'] for r in db.rows('SELECT id FROM jobs ORDER BY first_seen_at,id')]
        before=db.rows('SELECT COUNT(*) n FROM job_analyses')[0]['n'];processed=[];errors=[];paused=False
        for jid in ids:
            if should_stop and should_stop():paused=True;break
            if candidate_snapshot_hash(db)!=candidate['_cv_facts_hash']:raise RuntimeError('PROFESSIONAL_FACTS_CHANGED_RESTART_REQUIRED')
            try:
                result=analyze_job(db,jid,profile,candidate=candidate,include_out_of_scope=True)
                assert result['academic_review']['facts_hash']==candidate['_cv_facts_hash']
                processed.append(jid)
            except Exception as exc:
                code=type(exc).__name__;errors.append({'job_id':jid,'error_code':code})
                db.error(None,'ACADEMIC_REEVALUATION',code,job_id=jid)
        return {'generated_at':now(),'jobs_seen':len(ids),'jobs_processed':len(processed),'processed_job_ids':processed,
                'new_analyses':db.rows('SELECT COUNT(*) n FROM job_analyses')[0]['n']-before,
                'errors':errors,'paused':paused,'all_stored_jobs_processed':len(processed)==len(ids),
                'candidate_facts_hash':candidate['_cv_facts_hash'],'dam':dam_state(candidate),
                'external_requests':0,'submitted':0,'live_enabled':False}


def candidate_snapshot_hash(db):
    # No repeated DPAPI decrypts in the loop; use the exact same CV fingerprint.
    from .config import digest
    return digest([(r['id'],r['field']) for r in db.rows("SELECT id,field FROM candidate_facts WHERE revoked_at IS NULL AND field NOT LIKE 'answer:%' ORDER BY id")])


def repair_stored_ashby_modes(db):
    """Correct only documented workplace enums; preserve every old content snapshot."""
    from .config import digest
    from .geography import scope_for
    from .worker import worker_lock
    from .repository import unpack_job
    changed=[];profile=load_profile()
    with worker_lock(db),db.transaction() as con:
        con.execute('BEGIN IMMEDIATE')
        if db.rows("SELECT id FROM scan_runs WHERE status='running'"):raise RuntimeError('SEARCH_ALREADY_ACTIVE')
        rows=db.rows("SELECT j.* FROM jobs j WHERE EXISTS(SELECT 1 FROM job_sources js JOIN sources s ON s.id=js.source_id WHERE js.job_id=j.id AND s.kind='ashby')")
        for row in rows:
            raw=json.loads(row['raw']);mode={'Hybrid':'HYBRID','OnSite':'ONSITE','Remote':'REMOTE'}.get(raw.get('workplaceType'))
            if not mode or mode==row['mode']:continue
            data={**row,'mode':mode};new_hash=digest({k:data.get(k) for k in ('title','company','location','description','mode','apply_url')})
            checks=db.rows('SELECT * FROM commute_checks WHERE job_id=? ORDER BY id DESC',(row['id'],))
            scope,reason,_=scope_for(unpack_job(data),row['category'],profile,checks)
            # No fresh observation or verification is claimed by a local correction.
            db.execute("UPDATE jobs SET mode=?,content_hash=?,scope=?,scope_reason=?,analysis_status='PENDIENTE' WHERE id=? AND content_hash=?",(mode,new_hash,scope,reason,row['id'],row['content_hash']))
            db.execute('INSERT OR IGNORE INTO job_snapshots(job_id,content_hash,captured_at,description,raw) VALUES(?,?,?,?,?)',(row['id'],new_hash,now(),row['description'],row['raw']))
            change={'job_id':row['id'],'old_mode':row['mode'],'new_mode':mode,'old_hash':row['content_hash'],'new_hash':new_hash,'source_field':'workplaceType'}
            db.event('ASHBY_WORKPLACE_NORMALIZATION_CORRECTED',change,job_id=row['id']);changed.append(change)
    return changed


def opportunity_report(db):
    from .ui.model import load_view
    from .autoapply.selection_ui import reason_text
    from .autoapply.academic import LABELS
    view=load_view(db);offers=[]
    for job in view['jobs']:
        if job['it_status']!='CONFIRMED_IT':continue
        selection=job['selection'];academic=job['academic_review']
        pending=[i['reason'] for i in academic['items'] if i['status']=='PENDIENTE' and i['level']!='DESEABLE']
        offers.append({'job_id':job['id'],'company':job['company'],'title':job['title'],'location':job['location'],
            'url':job.get('apply_url') or job['url'],'open_status':job['open_status'],'last_verified_at':job['last_verified_at'],
            'deadline_status':job['deadline_status'],'deadline_at':job['deadline_at'],
            'reasonable_local_opportunity':job['potential'],'priority':job['review_priority'],'academic':job['academic_label'],
            'review_notes':job['priority_note'],
            'academic_pending':pending,'academic_items':academic['items'],'minimum_requested_years':job['minimum_years'],
            'experience':job['experience_label'],'distance':job['distance_label'],'transport':'Pendiente salvo evidencia de trayecto y turno vigente',
            'exclusions':[reason_text(c) for c in selection.get('exclusions',[])],
            'warnings':[reason_text(c) for c in selection.get('warnings',[])],
            'operational_pending':[reason_text(c) for c in selection.get('operational_pending',[])],
            'cv_approved':view['cv_approved'],'application':job['application'],'eligibility_for_submission':'HOLD: requiere todas las comprobaciones y autorizaciones; esta lista solo permite revisar/preparar.'})
    offers.sort(key=lambda r:(not r['reasonable_local_opportunity'],r['priority'],r.get('deadline_at') or '9999',r['company'],r['title']))
    pending=[{'job_id':j['id'],'company':j['company'],'title':j['title'],'location':j['location'],
        'url':j.get('apply_url') or j['url'],'academic':j['academic_label'],
        'reason':'HOLD: faltan funciones informáticas explícitas y/o ubicación. No preparar documentos ni automatizar hasta acreditar funciones IT.'} for j in view['jobs'] if j.get('pending_it_review')]
    return {'generated_at':now(),'all_jobs':len(view['jobs']),'it_jobs':len(offers),
            'reasonable_local_opportunities':sum(r['reasonable_local_opportunity'] for r in offers),
            'priority_local_opportunities':view['priority_local'],'conditional_local_opportunities':view['conditional_local'],
            'current_academic_analyses':view['academic_reviewed'],'cv_approved':view['cv_approved'],
            'offers':offers,'pending_it_reviews':pending,'live_enabled':False,'all_confirmed_it_offers_included':True,
            'limits':'Localidad Cork es una pista, no prueba de radio/recorrido. Dudas menores no excluyen; matrícula, idiomas, equivalencia, vigencia y transporte pendientes se muestran. No se preparan documentos sin CV aprobado.'}


def report_markdown(report):
    lines=['# Oportunidades IT tras terminar DAM','',
        'Generado: '+report['generated_at'],
        f"Examinadas: {report['all_jobs']}; IT: {report['it_jobs']}; locales razonables para revisar: {report['reasonable_local_opportunities']}.",'',report['limits'],'',
        'LIVE OFF. CV pendiente de aprobación.' if not report['cv_approved'] else 'LIVE OFF. CV aprobado por el usuario.','']
    for group,label in [(True,'Oportunidades locales para revisar'),(False,'Otras IT conservadas: remotos, fuera de área o excluidas')]:
        lines+=['## '+label,'']
        for r in report['offers']:
            if r['reasonable_local_opportunity']!=group:continue
            lines += ['### '+r['company']+' — '+r['title'],'',
                '[Anuncio oficial]('+r['url']+') · '+r['location']+' · '+r['open_status'],
                'Plazo: '+str(r['deadline_at'] or 'No publicado')+' · '+r['deadline_status']+'.',
                'Experiencia: '+r['experience']+'. '+r['distance']+'.',
                'Formación: '+r['academic']+'.',
                'Revisión: '+r['review_notes'],
                'Por contrastar: '+('; '.join(r['academic_pending']+r['warnings']+r['operational_pending']) or 'Revalidación final completa pendiente')+'.',
                'Exclusiones: '+('; '.join(r['exclusions']) or 'Ninguna acreditada en esta selección')+'.','']
    lines+=['## Anuncios con funciones IT pendientes, sin autorización para preparar','']
    for row in report.get('pending_it_reviews',[]):
        lines+=['- **'+row['company']+' — '+row['title']+'**: '+row['location']+'. '+row['reason']+' [Anuncio]('+row['url']+')']
    return '\n'.join(lines)


def report_excel(report):
    import io
    from openpyxl import Workbook
    from .exporting import safe_cell
    workbook=Workbook();ws=workbook.active;ws.title='Oportunidades IT'
    columns={'company':'Empresa','title':'Puesto','location':'Localidad','reasonable_local_opportunity':'Local para revisar',
        'priority':'Prioridad de revisión (0 primero)','open_status':'Vigencia','last_verified_at':'Última verificación','deadline_status':'Estado plazo','deadline_at':'Fecha límite publicada',
        'academic':'Formación','academic_pending':'Dudas académicas','review_notes':'Condiciones de revisión','experience':'Experiencia mínima publicada',
        'distance':'Distancia verificada','warnings':'Advertencias','exclusions':'Exclusiones',
        'operational_pending':'Pendientes operativos','application':'Estado candidatura','url':'Enlace oficial'}
    ws.append(list(columns.values()))
    for row in report['offers']:
        values=[]
        for key in columns:
            value=row.get(key)
            if isinstance(value,list):value='; '.join(value)
            values.append(safe_cell(value))
        ws.append(values)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=35
    info=workbook.create_sheet('Alcance');info.append(['Generado',report['generated_at']]);info.append(['Ofertas examinadas',report['all_jobs']])
    info.append(['IT',report['it_jobs']]);info.append(['Locales para revisar',report['reasonable_local_opportunities']]);info.append(['Límites',report['limits']])
    info.append(['CV aprobado',report['cv_approved']]);info.append(['LIVE',False])
    pending=workbook.create_sheet('IT por confirmar');pending.append(['Empresa','Puesto','Ubicación','Formación','Motivo HOLD','URL'])
    for row in report.get('pending_it_reviews',[]):pending.append([safe_cell(row[k]) for k in ('company','title','location','academic','reason','url')])
    buffer=io.BytesIO();workbook.save(buffer);return buffer.getvalue()
