import csv
import io
import json
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from .config import ROOT, now
from .repository import all_jobs,latest_analyses,run_observations
from .audit import coverage, coverage_markdown

LABELS={'id':'ID','title':'Título','company':'Empresa','location':'Ubicación publicada','url':'Enlace original',
        'apply_url':'Enlace de candidatura','description':'Descripción original','detail_complete':'Texto íntegro recuperado',
        'category':'Sector','scope':'Alcance','scope_reason':'Motivo del alcance','mode':'Modalidad','posted_at':'Fecha publicada',
        'first_seen_at':'Primera observación','last_seen_at':'Última observación','last_verified_at':'Última verificación',
        'open_status':'Vigencia','analysis_status':'Estado del análisis','content_hash':'Hash de contenido','raw':'Datos originales',
        'possible_duplicate':'Posible duplicado','sources':'Fuentes','fit_score':'Encaje provisional','data_confidence':'Confianza de datos',
        'depth':'Profundidad','salary':'Salario publicado','schedule':'Horario publicado','contract':'Contrato publicado','experience':'Experiencia publicada',
        'source_id':'ID fuente','job_id':'ID oferta','run_id':'ID ejecución','name':'Nombre','kind':'Tipo','config':'Configuración',
        'enabled':'Activada','permission':'Método permitido','status':'Estado','checked_at':'Fecha de comprobación','last_success':'Último éxito',
        'notes':'Notas','pages_visited':'Páginas recorridas','expected_total':'Total anunciado','observed':'Publicaciones observadas',
        'pagination_complete':'Paginación completa','cursor':'Cursor','message':'Mensaje','started_at':'Inicio','finished_at':'Fin',
        'occurred_at':'Fecha de incidencia','created_at':'Fecha de registro','details':'Detalle','action':'Marca personal',
        'application_date':'Fecha de candidatura','external_id':'ID de origen','captured_at':'Fecha de captura','result':'Resultado',
        'profile_hash':'Versión del perfil','analyzer_version':'Versión del analizador','evidence':'Evidencia','generated_at':'Fecha de generación',
        'change_status':'Cambio en ejecución','analysis_stale':'Análisis desactualizado','summary_es':'Resumen en español',
        'verdict_es':'Recomendación razonada','fit_range':'Intervalo por incertidumbre','weekly_hours':'Jornada semanal',
        'duties_original':'Funciones publicadas','mandatory_original':'Requisitos obligatorios publicados','desirable_original':'Requisitos deseables publicados',
        'semantic_es':'Interpretaciones locales en español con citas','missing_fields':'Datos no acreditados','risks':'Dificultades','application_status':'Estado de candidatura',
        'score_explanation':'Desglose del encaje','profile_evidence':'Contraste con competencias confirmadas'}

def safe_cell(value):
    if isinstance(value,(dict,list)):
        value = json.dumps(value,ensure_ascii=False)
    if isinstance(value,str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',value)
        if value.lstrip().startswith(('=','+','-','@')):
            value = "'"+value
    return value

def export_all(db, output_dir=None, run_id=None):
    with db.read_snapshot() as snapshot:
        return _export_snapshot(snapshot,output_dir,run_id)

def _export_snapshot(db, output_dir=None, run_id=None):
    output_dir = Path(output_dir or ROOT/'exports')
    output_dir.mkdir(parents=True,exist_ok=True)
    report = coverage(db,run_id)
    analyses=latest_analyses(db)
    rows = all_jobs(db,analyses)
    observations=run_observations(db,report['run_id']) if report['run_id'] else {}
    for row in rows:
        result=analyses.get(row['id'],{})
        row['change_status']=observations.get(row['id'],'NO_OBSERVADA')
        for key in ('summary_es','verdict_es','fit_range','missing_fields','risks','score_explanation','profile_evidence'):
            row[key]=result.get(key)
        for key in ('duties_original','mandatory_original','desirable_original'):
            row[key]=result.get('documentary',{}).get(key)
        row['weekly_hours']=result.get('facts',{}).get('weekly_hours',{}).get('value')
        row['semantic_es']=result.get('semantic',{}).get('items')
    metadata = {'generated_at':now(),'run_id':report['run_id'],'count':len(rows),'note':'Todas las ofertas locales; auditoría de la ejecución indicada'}
    stamp = metadata['generated_at'].replace(':','').replace('+','_')
    base = output_dir / f'cork_job_radar_{stamp}'
    tables = {
        'Todas_las_ofertas':rows,
        'Nuevas':[r for r in rows if r['change_status']=='NEW'],
        'Cambiadas':[r for r in rows if r['change_status']=='CHANGED'],
        'Repetidas':[r for r in rows if r['change_status']=='REPEATED'],
        'Recomendadas':[r for r in rows if (r.get('fit_score') or 0)>=70 and r['scope']=='EN_ALCANCE' and r['open_status'] not in {'CLOSED_CONFIRMED','UNLISTED'}],
        **{cat:[r for r in rows if r['category']==cat] for cat in ('IT','Produccion','Almacen','Tecnicos','Otras')},
        'Pendientes':[r for r in rows if r['analysis_status']!='COMPLETO'],
        'Accesibilidad_pendiente':[r for r in rows if r['scope']=='ACCESIBILIDAD_PENDIENTE'],
        'Fuera_de_alcance':[r for r in rows if r['scope']=='FUERA_DE_ALCANCE'],
        'Cerradas':[r for r in rows if r['open_status']=='CLOSED_CONFIRMED'],
        'Fuentes':db.rows('SELECT * FROM sources'), 'Auditoria':report['source_runs'],
        'Busquedas':db.rows('SELECT * FROM scan_runs'),
        'Observaciones':db.rows('SELECT * FROM run_jobs'),
        'Descubrimientos':db.rows('SELECT * FROM discovery_runs'),
        'Recursos_descubrimiento':db.rows('SELECT * FROM discovery_resources'),
        'Enlaces_por_revisar':db.rows('SELECT * FROM discovery_candidates'),
        'Errores':db.rows('SELECT * FROM source_errors'),
        'Historial':db.rows('SELECT * FROM audit_events'),
        'Candidaturas':db.rows('SELECT * FROM user_actions'),
        'AutoApply_planes':db.rows('SELECT * FROM application_plans'),
        'AutoApply_intentos':db.rows('SELECT * FROM application_attempts'),
        'AutoApply_eventos':db.rows('SELECT * FROM application_events'),
        'AutoApply_cupos':db.rows('SELECT * FROM application_reservations'),
        'Procedencias':db.rows('SELECT * FROM job_sources'),
        'Verificaciones':db.rows('SELECT * FROM verification_events'),
        'Trayectos':db.rows('SELECT * FROM commute_checks'),
        'Analisis':db.rows('SELECT * FROM job_analyses'),
        'Snapshots':db.rows('SELECT * FROM job_snapshots'),
        'Metadatos':[metadata],
    }
    wb = Workbook()
    wb.remove(wb.active)
    long_text = []
    for name,data in tables.items():
        ws = wb.create_sheet(name)
        ws.append(['Generado',metadata['generated_at'],'Ejecución',metadata['run_id']])
        headers = list(data[0]) if data else ['Sin registros']
        ws.append([LABELS.get(k,k) for k in headers])
        for row_number,row in enumerate(data,3):
            values = []
            for key in headers:
                value = safe_cell(row.get(key))
                if isinstance(value,str) and len(value)>32000:
                    for part in range(0,len(value),30000):
                        long_text.append({'sheet':name,'row':row_number,'field':key,'part':part//30000+1,'text':value[part:part+30000]})
                    value = value[:31000]+'\n[Texto íntegro por partes en Textos_largos]'
                values.append(value)
            ws.append(values)
            for index,key in enumerate(headers,1):
                if key in {'url','apply_url'} and isinstance(row.get(key),str) and row[key].startswith(('http://','https://')):
                    ws.cell(row_number,index).hyperlink=row[key]
        ws.freeze_panes = 'A3'
        ws.auto_filter.ref = f'A2:{get_column_letter(len(headers))}{max(2,ws.max_row)}'
        for cell in ws[2]:
            cell.font = Font(bold=True)
        for index,key in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(index)].width = min(55,max(15,len(key)+3))
    if long_text:
        ws = wb.create_sheet('Textos_largos')
        ws.append(list(long_text[0]))
        for row in long_text:
            ws.append([safe_cell(v) for v in row.values()])
    xlsx = base.with_suffix('.xlsx')
    wb.save(xlsx)
    with base.with_suffix('.csv').open('w',encoding='utf-8-sig',newline='') as stream:
        writer = csv.DictWriter(stream,fieldnames=[LABELS.get(k,k) for k in rows[0]] if rows else ['Sin ofertas'],delimiter=';')
        writer.writeheader()
        writer.writerows({LABELS.get(k,k):safe_cell(v) for k,v in row.items()} for row in rows)
    payload = {'metadata':metadata,'coverage':report,'tables':tables}
    base.with_suffix('.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    base.with_suffix('.md').write_text(coverage_markdown(report),encoding='utf-8')
    return {'xlsx':str(xlsx),'csv':str(base.with_suffix('.csv')),'json':str(base.with_suffix('.json')),
            'coverage':str(base.with_suffix('.md')),'count':len(rows),'metadata':metadata,
            'empty_jobs_notice':'Sin ofertas; el archivo solo contiene auditoría/configuración real' if not rows else None}
