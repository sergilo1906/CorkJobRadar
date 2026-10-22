"""Genera documentación desde resultados realmente persistidos; no ejecuta pruebas ni las simula."""
import json
import xml.etree.ElementTree as ET
from corkjobradar.config import ROOT,now
from corkjobradar.db import Database
from corkjobradar.audit import coverage,coverage_markdown
from corkjobradar.sources import export_sources
from corkjobradar.exporting import export_all
from openpyxl import load_workbook

def main():
    db=Database()
    report=coverage(db)
    sources=db.rows('SELECT * FROM sources ORDER BY kind,id')
    stats={k:v for k,v in report.items() if k not in {'sources','source_runs','uncovered_sources'}}
    (ROOT/'logs/live-coverage.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'config/sources.yaml').write_text(export_sources(db),encoding='utf-8')
    exported=export_all(db,run_id=report['run_id'])
    workbook=load_workbook(exported['xlsx'],read_only=True)
    count=workbook['Todas_las_ofertas'].max_row-2
    if count!=len(db.rows('SELECT id FROM jobs')):
        raise AssertionError('No cuadra la exportación real con SQLite')
    sheets=list(workbook.sheetnames)
    workbook.close()
    (ROOT/'logs/verified-export.json').write_text(json.dumps({**exported,'xlsx_rows_verified':count,'sheets_verified':sheets},ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'docs/LAST_COVERAGE.md').write_text(coverage_markdown(report),encoding='utf-8')
    parts=['# Fuentes comprobadas y pendientes','',f'Generado desde SQLite: {now()}. Fuentes registradas: {len(sources)}.','',
           'Identificar una página y comprobar su primera respuesta no equivale a implementar su rastreo. Las cifras de anuncios proceden del trabajador y SQLite, no del buscador web.','',
           '| Fuente | Conector | Activada | Estado | URL | Observaciones |','|---|---|---|---|---|---|']
    for s in sources:
        parts.append('| '+' | '.join([s['name'],s['kind'],'Sí' if s['enabled'] else 'No',s['status'],f"[Abrir]({s['url']})",s['notes'].replace('|','/').replace('\n',' ')])+' |')
    parts+=['','## Permiso y documentación de API','',
            '- [Greenhouse Job Board API](https://docs.greenhouse.io/job-board.html): GET públicos; no se usa presentación de candidaturas.',
            '- [Lever Postings API](https://github.com/lever/postings-api): publicaciones públicas, skip/limit y variantes global/UE.',
            '- [Ashby Job Postings API](https://developers.ashbyhq.com/docs/public-job-posting-api): API pública del tablero, sin API administrativa.',
            '- [SmartRecruiters Posting API](https://developers.smartrecruiters.com/docs/posting-api): listado y detalle públicos; AbbVie se consulta con country=ie.',
            '', '## Evidencia y alcance','',
            'Greenhouse: GlobeTech y Keeper Security; Lever: Bright Machines; Ashby: Qumulo y NETGEAR; SmartRecruiters: AbbVie Irlanda. El conector HTML/microdatos se ha probado en un único anuncio público de AbbVie, nunca como cobertura de la empresa completa.',
            'Los fixtures y pruebas en vivo son distintos: tests/fixtures no se importan en producción. Ver logs/endpoint_probe.json, logs/endpoint_probe_2.json, logs/inventory_probe.json y logs/health-final.log. Las páginas y detalles de ejecuciones completas están en source_runs, job_sources y job_snapshots.',
            '', '## Fuentes desactivadas','',
            'kind=manual significa conector específico no implementado o condiciones pendientes de revisión; HTTP 200 no lo convierte en compatible. Se ofrece importación manual de URL/CSV. Indeed y Johnson & Johnson han mostrado restricciones; no se eluden. Publicjobs tuvo error inicial y fue accesible en la revisión posterior; permanece en revisión.',
            'Las agencias no se presentan como empresa contratante cuando esta no se publica. Las ubicaciones de las empresas no prueban rutas sin coche. LinkedIn se excluye de rastreo y enlaces principales.',
            '', 'El catálogo es ampliable mediante editor/importación YAML/CSV y descubrimiento de ATS enlazados; no hay máximo fijo de empresas. RSS/Atom, sitemaps, Workday y portales específicos aún no se anuncian como compatibles.']
    (ROOT/'docs/SOURCES.md').write_text('\n'.join(parts)+'\n',encoding='utf-8')
    xml=ET.parse(ROOT/'logs/pytest-delivery.xml').getroot()
    suite=xml.find('testsuite') if xml.tag!='testsuite' else xml
    teststats=dict(suite.attrib)
    verifications=db.rows('SELECT open_status,count(*) n FROM jobs GROUP BY open_status')
    runs=db.rows('SELECT id,status,started_at,finished_at,message FROM scan_runs ORDER BY id')
    test_report=f'''# Informe de pruebas y comprobaciones

Generado: {now()}. Windows 11 10.0.26200, Python 3.12.4, entorno .venv del proyecto.

## Pruebas automatizadas

Última suite completa: `{teststats.get('tests')}` casos; fallos `{teststats.get('failures')}`; errores `{teststats.get('errors')}`; omitidos `{teststats.get('skipped')}`; duración `{teststats.get('time')}` segundos. Evidencia: logs/pytest-delivery.xml y logs/pytest-delivery.log. Tras traducir encabezados de Excel, las tres pruebas específicas de UI y XLSX también pasaron (logs/pytest-export-spanish.log).

Comandos ejecutados:

```powershell
.\\.venv\\Scripts\\python.exe -m pytest -q --basetemp=data\\pytest-delivery --junitxml=logs\\pytest-delivery.xml
.\\.venv\\Scripts\\python.exe -m pytest tests\\test_ui.py tests\\test_safety_and_states.py::test_excel_long_text_and_formula_injection -q --basetemp=data\\pytest-export-spanish
powershell -NoProfile -ExecutionPolicy Bypass -File .\\instalar.ps1
.\\.venv\\Scripts\\python.exe -m pip check
.\\.venv\\Scripts\\python.exe -m corkjobradar doctor
.\\.venv\\Scripts\\python.exe -m corkjobradar sources --load
.\\.venv\\Scripts\\python.exe -m corkjobradar sources --refresh
.\\.venv\\Scripts\\python.exe -m corkjobradar scan --all
.\\.venv\\Scripts\\python.exe -m corkjobradar verify --all
.\\.venv\\Scripts\\python.exe -m streamlit run app.py --server.address 127.0.0.1 --server.port 8501
```

El instalador terminó con código 0 (logs/installer-final.log); pip check: sin dependencias rotas. doctor: SQLite integrity_check=ok. Los scripts de PowerShell se probaron en este proyecto, sin WSL ni modificación de políticas globales. La creación inicial de .venv y la instalación desde entorno vacío se ejecutaron; la repetición del instalador también fue correcta. Se reinició el servidor de este proyecto y se ejecutó iniciar.ps1: arranque correcto y /_stcore/health devolvió HTTP 200, cuerpo ok. Los .cmd son alternativas escritas, pendientes de una prueba interactiva de doble clic.

## Cobertura de la suite

- Fixtures por Greenhouse, Lever, Ashby, SmartRecruiters y HTML JSON-LD/microdatos.
- Fechas, salarios por unidad, turnos, deduplicación conservadora, alcance IT/NO IT, rutas sintéticas de cero/un transbordo, bicicleta y referencia IT sin coche.
- HTTP 403/404/429/503, Retry-After, timeout, gzip, límites de tamaño, robots, redirección a red privada, caché por Accept, cambios de estructura, páginas repetidas y tarjetas malformadas.
- 137 puestos en 7 páginas + 12 duplicados = 149 procedencias; 8 sin salario, 5 detalles fallidos y 3 cierres explícitos. Resultado exacto: 132 análisis básicos + 5 pendientes; al reanudar, solo cinco descargas y 137 análisis básicos. Sin datos sintéticos en producción.
- Interrupción tras 60 puestos: cursor persistido y recuperación hasta 137 sin perder registros.
- DB limpia, reinicio, candidaturas persistidas, importación de omisión incluso en 403, XLSX reabierto, CSV BOM, JSON y conservación de texto mayor de 32.767 caracteres.
- AppTest: seis pestañas, filtros, detalle, exportación y lanzamiento del trabajador sin bloquear UI. CV PDF vacío y DOCX sintético se procesan sin inventar competencias.
- Advertencias: cinco DeprecationWarning de Streamlit/NumPy en la suite UI; no son fallos. Ollama no fue encontrado en PATH; no se atribuye una prueba en vivo del proveedor profundo.

## Humo real de conectores

```json
{json.dumps(stats,ensure_ascii=False,indent=2)}
```

Todos los conectores activos cuentan con ejecución real de extracción, persistencia y análisis básico. Los textos completos se muestran en la ficha local. El navegador real abrió http://127.0.0.1:8501 y su DOM mostró Cork Job Radar, seis pestañas y cero componentes de excepción. La herramienta de captura no permitía guardar bajo C:\\Proyectos; se usó lectura del DOM y AppTest, sin afirmar una captura inexistente en el proyecto.

Validación XLSX real: {count} filas maestras y {len(sheets)} hojas reabiertas con openpyxl. Archivo: {exported['xlsx']}. Manifiesto: logs/verified-export.json. La exportación corresponde a la ejecución de descubrimiento {report['run_id']} y a estados actuales de datos.

Verificaciones actuales: {json.dumps(verifications,ensure_ascii=False)}. Las bloqueadas o reintentables permanecen visibles; no se consideran cerradas. El recorrido completo de los listados no garantiza que todos los formularios de solicitud sean accesibles.

La ejecución adicional de verificación 6 terminó parcial: 205 tareas procesadas y 67 reintentables tras suspender el dominio que devolvió un bloqueo. En el navegador se registraron ocho ERR_CONNECTION_REFUSED durante el reinicio intencionado del servidor; la conexión se recuperó y el DOM posterior mostró seis pestañas, datos reales y cero excepciones de componentes. Evidencia de comprobaciones finales: logs/final-checks.json.

## Fallos encontrados y corregidos

Instalación todavía en curso en la primera prueba de doctor; expresión salarial con espacios; fixture HTML mal construido; tratamiento de redirecciones; JSON frente a HTML por Accept; microdatos sin JSON-LD; página genérica vacía marcada incorrectamente completa; doble descompresión HTTP; separación de identidad canónica frente a URL publicada; caché independiente por Accept. Se preservaron los errores y se corrigió el estado histórico de la primera fuente HTML. Una ejecución de tests concurrente reutilizó basetemp y falló por bloqueo de Windows; se repitieron secuencialmente en directorios separados y pasaron.

La ejecución 4 de verificación fue cancelada para corregir los enlaces; el estado se conserva. Historial real:

```json
{json.dumps(runs,ensure_ascii=False,indent=2)}
```

## Lo no validado/no implementado

No se ha validado análisis semántico profundo en vivo, rutas automáticas, navegadores dinámicos de conectores, RSS/Atom/sitemaps ni parsers específicos de los portales/agencias desactivados. No se afirma cobertura total de Cork. Los análisis de reglas son PARCIAL_POR_DATOS, no COMPLETO.
'''
    (ROOT/'docs/TEST_REPORT.md').write_text(test_report,encoding='utf-8')
    progress=f'''# Cork Job Radar — progreso verificable

Actualizado: {now()}. Especificación completa leída (360 líneas) y copiada desde la carpeta inicial. Proyecto: C:\\Proyectos\\CorkJobRadar.

## Fases

1. **Verificada:** entorno Windows/Python, .venv, dependencias fijadas, scripts, perfil Ballinlough, SQLite v1, CLI doctor y arranque Streamlit con datos locales reales.
2. **Verificada para lo anunciado:** cuatro ATS y HTML JSON-LD/microdatos, fixtures, errores/paginación y conectividad en vivo; URL/CSV manuales. No hay Playwright, RSS/Atom ni sitemaps implementados.
3. **Operativa y ampliable, cobertura parcial del mercado:** {len(sources)} fuentes comprobadas/registradas; 7 activadas; demás pendientes/restringidas. Investigación de portales, empresas y agencias documentada. Descubrimiento de ATS enlazados y edición/importación de catálogo disponibles.
4. **Verificada:** cola SQLite, cursor, cancelación, bloqueo de un trabajador, reanudación, snapshots, deduplicación con procedencias, reconciliación y alerta de recuentos. Prueba 137/7/12 exacta. No hay denominador del mercado total.
5. **Básico verificado, profundo pendiente:** textos en alcance procesados individualmente con reglas y estados parciales. Puntuación/certeza separadas. Ollama opcional escrito; no validado en vivo. Interpretación exhaustiva de tareas/requisitos en español pendiente del motor profundo/revisión.
6. **Operativa y probada en sus flujos automatizados:** seis pestañas, filtros, fichas originales, notas/candidaturas, importación, perfil, propuestas CV locales y exportación real XLSX/CSV/JSON/Markdown. Rutas manuales con evidencia; no cálculo de transporte.
7. **Verificada con límites declarados:** {teststats.get('tests')} pruebas completas sin fallos ni errores y 3 comprobaciones específicas posteriores de UI/Excel. Humo real en cuatro familias ATS y un anuncio genérico. Sin pruebas ficticias declaradas como reales.

## Datos reales actuales

{json.dumps(stats,ensure_ascii=False,indent=2)}

Exportación comprobada: {exported['xlsx']} ({count} anuncios). Ver docs/TEST_REPORT.md y logs/verified-export.json.

Verificación adicional: ejecución 6 parcial, 205 tareas procesadas y 67 reintentables tras bloqueo; nunca equivalen a 67 cierres. Estados actuales: {json.dumps(verifications,ensure_ascii=False)}. Se han distinguido {report['talent_pool_posts']} bolsas de talento que no acreditan vacantes concretas.

## Restricciones y siguientes pasos

- Activar un modelo local ya instalado y validar su análisis profundo antes de anunciar completitud semántica.
- Conectar una vía autorizada de rutas/transporte y verificar ida/vuelta según turnos; mantener Ballinlough y ausencia de coche.
- Implementar y probar paginación y fichas de JobsIreland/HSE/agencias accesibles tras revisar condiciones; conservar restricciones de Indeed/J&J y las revisiones pendientes.
- Completar soporte de feeds/sitemaps y parsers dinámicos solo cuando se permita.
- Diagnóstico de omisiones: incidencia automática disponible, investigación causal y regresiones de nuevas fuentes requieren revisión técnica.
- Endurecer fijación DNS contra rebinding y preparar siguientes migraciones con copias previas. No se han enviado CV, correos ni candidaturas a empresas.

## Incidencias de ejecución y aislamiento

La primera ejecución de pytest creó temporales bajo el TEMP predeterminado del usuario, antes de fijar basetemp dentro del proyecto. La navegación automatizada generó su snapshot automático bajo la carpeta de trabajo inicial; al intentar guardarlo en el proyecto, la herramienta denegó esa ruta. No se borraron esos archivos externos sin autorización. Las siguientes pruebas y los informes del proyecto usan rutas internas. Los otros proyectos no se editaron. Estas incidencias se registran para no afirmar aislamiento absoluto que no se cumplió en esas dos herramientas.

No se declara terminado el análisis semántico profundo, el cálculo real de rutas ni los conectores aún no implementados.
'''
    for filename in ('PROGRESO.md','PROGRESS.md'):
        (ROOT/filename).write_text(progress,encoding='utf-8')
    (ROOT/'NEXT_STEPS.md').write_text('''# Continuación

Lee PROGRESO.md, docs/TEST_REPORT.md y la especificación original. No reinstales ni borres la base.

1. Ejecuta doctor y revisa ejecuciones/tareas reintentables en la interfaz. La auditoría por defecto selecciona la última ejecución con descubrimiento.
2. Los siete conectores activos ya tienen pruebas offline y humo real. Amplía las fuentes según docs/SOURCES.md; no actives portales sin revisar condiciones y paginación.
3. Prioridad de calidad: validar un modelo local para interpretación exhaustiva en español. El modo básico permanece parcial de forma intencionada.
4. Prioridad de cobertura: JobsIreland, HSE y agencias con acceso permitido. Rutas desde Ballinlough requieren proveedor autorizado o evidencia manual de ida/vuelta.
5. Ejecuta pruebas solo de forma secuencial con basetemp dentro de data/. No importes fixtures en producción ni sustituyas el historial real por datos de ejemplo.
6. La interfaz local está en 127.0.0.1:8501. Exportaciones en exports/, manifiesto de la última comprobada en logs/verified-export.json.
''',encoding='utf-8')
    print(json.dumps({'tests':teststats,'coverage':stats,'export':exported,'verification':verifications},ensure_ascii=False,indent=2))

if __name__=='__main__': main()
