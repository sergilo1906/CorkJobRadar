"""Audit real persisted data and reopen an actual export. No application sends."""
from pathlib import Path
import hashlib
import json
import sqlite3
import sys
from xml.etree import ElementTree as ET

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from corkjobradar.config import now
from corkjobradar.db import Database
from corkjobradar.autoapply.engine import report
from corkjobradar.exporting import export_all
from openpyxl import load_workbook


def main():
    baseline=json.loads((ROOT/'logs/autoapply-diagnostic-initial.json').read_text(encoding='utf-8'))
    backup=Path(baseline['databases']['radar.sqlite3']['path'])
    db=Database()
    retention={}
    with sqlite3.connect(backup.as_uri()+'?mode=ro',uri=True) as old,sqlite3.connect(db.path.as_uri()+'?mode=ro',uri=True) as current:
        current.execute('BEGIN')
        for table in baseline['databases']['radar.sqlite3']['counts']:
            before=old.execute('SELECT * FROM "'+table+'" ORDER BY rowid').fetchall()
            after=current.execute('SELECT * FROM "'+table+'" ORDER BY rowid').fetchall()
            retention[table]={'before':len(before),'after':len(after),'all_rows_equal':before==after}
            if table!='http_cache':assert before==after,table+' changed historical data'
        assert current.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert current.execute('PRAGMA foreign_key_check').fetchall()==[]
        assert current.execute('SELECT COUNT(*) FROM user_actions').fetchone()[0]==0
        assert current.execute('SELECT COUNT(*) FROM application_attempts').fetchone()[0]==0
        assert current.execute('SELECT COUNT(*) FROM candidate_facts').fetchone()[0]==0
        assert current.execute("SELECT COUNT(*) FROM jobs WHERE json_extract(raw,'$._synthetic_fixture')=1").fetchone()[0]==0
    exported=export_all(db,ROOT/'exports/autoapply-verified')
    workbook=load_workbook(exported['xlsx'],read_only=True,data_only=False)
    sheets={name:workbook[name].max_row-2 for name in ('Todas_las_ofertas','AutoApply_planes','AutoApply_intentos','AutoApply_eventos','AutoApply_cupos','Candidaturas')}
    workbook.close()
    expected=baseline['databases']['radar.sqlite3']['counts']['jobs']
    assert sheets['Todas_las_ofertas']==expected and sheets['AutoApply_planes']==expected
    assert sheets['AutoApply_intentos']==sheets['Candidaturas']==0
    payload=json.loads(Path(exported['json']).read_text(encoding='utf-8'))
    assert len(payload['tables']['AutoApply_planes'])==expected
    assert len(payload['tables']['Todas_las_ofertas'])==expected
    assert Path(exported['csv']).read_bytes().startswith(b'\xef\xbb\xbf')
    test_results={}
    for path in sorted((ROOT/'logs').glob('autoapply-*.xml')):
        root=ET.parse(path).getroot();suites=list(root) if root.tag=='testsuites' else [root]
        test_results[path.name]={key:sum(float(s.get(key,0)) for s in suites) for key in ('tests','failures','errors','skipped','time')}
    result={'checked_at':now(),'status':report(db),'historical_retention':retention,'integrity':'ok','foreign_keys':[],
            'cache_note':'Public GET form inspection may add/update HTTP cache; original cache preserved in backup.',
            'export':exported,'reopened_excel_rows':sheets,'json_reopened':True,'csv_bom':True,
            'tests':test_results,'real_submissions':0,'synthetic_jobs_in_production':0,'candidate_facts_in_production':0}
    (ROOT/'logs/autoapply-final-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'integrity':'ok','all_historical_business_rows_equal':True,'sheets':sheets,'status':result['status']},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
