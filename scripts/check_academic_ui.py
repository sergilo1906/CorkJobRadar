"""Read-only actual Chromium proof: new education, complete export and unchanged CV consent."""
import io
import argparse
import json
import os
from pathlib import Path
import sys
from urllib.parse import urlsplit
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from corkjobradar.config import ROOT,now
from corkjobradar.db import Database
from corkjobradar.ui.model import load_view
from scripts.check_ui_redesign import navigate,idle,layout,select


def main():
    from playwright.sync_api import sync_playwright,expect
    from openpyxl import load_workbook
    parser=argparse.ArgumentParser();parser.add_argument('--label',default='academic-reevaluation');args=parser.parse_args()
    if not args.label.replace('-','').isalnum():raise ValueError('SAFE_LABEL_REQUIRED')
    output=ROOT/'logs'/(args.label+'-browser.json')
    if output.exists():raise ValueError('PRESERVE_EXISTING_PROOF')
    folder=ROOT/'data/ui-redesign'/args.label;folder.mkdir(exist_ok=False)
    os.environ['PLAYWRIGHT_BROWSERS_PATH']=str(ROOT/'data/playwright-browsers')
    os.environ['TEMP']=os.environ['TMP']=str(ROOT/'data/browser-temp')
    db=Database();view=load_view(db);expected_count=sum(j['it_status']=='CONFIRMED_IT' for j in view['jobs'])
    before={t:db.rows('SELECT * FROM '+t) for t in ('candidate_facts','candidate_documents','application_attempts','autoapply_control')}
    refused=[];screens=[];download_rows=0
    with sync_playwright() as pw:
        browser=pw.chromium.launch(args=['--disable-background-networking'])
        context=browser.new_context(viewport={'width':1440,'height':900},service_workers='block',accept_downloads=True)
        def allowed(url):
            a=urlsplit(url);return a.hostname=='127.0.0.1' and a.port==8501
        def guard(route):
            if allowed(route.request.url):route.continue_()
            else:refused.append(True);route.abort()
        context.route('**/*',guard);context.route_web_socket('**/*',lambda ws:ws.connect_to_server() if allowed(ws.url) else ws.close())
        page=context.new_page();page.goto('http://127.0.0.1:8501')
        page.get_by_role('heading',name='Tu próxima oportunidad IT',exact=True).wait_for(timeout=60000);idle(page)
        for width,height in [(1440,900),(390,844)]:
            page.set_viewport_size({'width':width,'height':height});page.wait_for_timeout(500)
            navigate(page,'Ofertas de empleo')
            expect(page.get_by_text('Reevaluación académica vigente:',exact=False)).to_be_visible()
            page.get_by_text('Lista completa tras actualizar mi formación',exact=True).click();idle(page)
            with page.expect_download() as event:page.get_by_role('button',name='Descargar lista IT en Excel',exact=True).click()
            path=folder/f'opportunities-{width}.xlsx';event.value.save_as(path)
            book=load_workbook(path);download_rows=book['Oportunidades IT'].max_row-1
            assert download_rows==expected_count
            assert book['IT por confirmar'].max_row-1==sum(j.get('pending_it_review',False) for j in view['jobs'])
            assert not page.get_by_test_id('stException').count()
            dims=layout(page);assert dims['document_width']<=width+2 and dims['overflow_controls']==0
            page.screenshot(path=str(folder/f'offers-{width}.png'),full_page=True);screens.append({'screen':'offers','width':width,'layout':dims})
            select(page,'Mostrar','IT por confirmar')
            assert page.get_by_role('button',name='Ver oferta',exact=True).count()==sum(j.get('pending_it_review',False) for j in view['jobs'])
            page.screenshot(path=str(folder/f'pending-it-{width}.png'),full_page=True)
            select(page,'Mostrar','Oportunidades IT')
            page.get_by_role('button',name='Ver oferta',exact=True).first.click();idle(page)
            expect(page.get_by_role('heading',name='Formación y titulación',exact=True)).to_be_visible()
            assert 'DAM terminado' in page.get_by_test_id('stMain').inner_text()
            page.screenshot(path=str(folder/f'job-{width}.png'),full_page=True)
            navigate(page,'Mis candidaturas')
            expect(page.get_by_text('Reevaluación académica vigente:',exact=False)).to_be_visible()
            assert page.get_by_role('button',name='Abrir candidatura',exact=True).count()>0
            page.screenshot(path=str(folder/f'applications-{width}.png'),full_page=True)
            navigate(page,'Perfil y configuración')
            expect(page.get_by_role('button',name='Aprobar esta versión del CV',exact=True)).to_be_disabled()
            expect(page.get_by_role('checkbox',name='He revisado íntegramente esta versión y apruebo este CV definitivo',exact=True)).not_to_be_checked()
        context.close();browser.close()
    for table,rows in before.items():assert db.rows('SELECT * FROM '+table)==rows,table
    assert not refused
    result={'at':now(),'status':'PASS','screens':screens,'excel_it_rows':download_rows,'academic_profile_visible':True,
        'application_opportunities_visible':True,'cv_version_2_pending_and_consent_unchecked':True,'unchanged_professional_document_send_tables':True,
        'external_http_requests':len(refused),'live_enabled':False}
    output.write_text(json.dumps(result,indent=2),encoding='utf-8');print(json.dumps(result))


if __name__=='__main__':main()
