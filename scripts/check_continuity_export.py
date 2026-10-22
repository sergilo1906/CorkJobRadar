"""Generate real exports from a consistent SQLite snapshot and cross-check their IDs."""
import csv
import hashlib
import json
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from corkjobradar.config import ROOT,now
from corkjobradar.db import Database
from corkjobradar.exporting import export_all


def main():
    result=export_all(Database())
    payload=json.loads(Path(result['json']).read_text(encoding='utf-8'))
    jobs=payload['tables']['Todas_las_ofertas'];expected={r['id'] for r in jobs}
    with Path(result['csv']).open(encoding='utf-8-sig',newline='') as stream:
        csv_rows=list(csv.DictReader(stream,delimiter=';'))
    workbook=load_workbook(result['xlsx'],read_only=True,data_only=False)
    try:
        sheet=workbook['Todas_las_ofertas'];headers=next(sheet.iter_rows(min_row=2,max_row=2,values_only=True));column=headers.index('ID')
        excel_ids=[];formulas=0
        for row in sheet.iter_rows(min_row=3):
            excel_ids.append(row[column].value);formulas+=sum(cell.data_type=='f' for cell in row)
        assert len(expected)==len(jobs)==result['count']==len(csv_rows)==len(excel_ids)
        assert set(excel_ids)=={r['ID'] for r in csv_rows}==expected
        assert formulas==0
        report={'checked_at':now(),'real_snapshot':True,'result':result,'unique_ids':len(expected),
                'xlsx_csv_json_ids_match':True,'formulas_in_job_sheet':formulas,'sheets':workbook.sheetnames,
                'global_inventory':payload['coverage']['global_analysis_inventory'],'live_submission_enabled':False,
                'sha256':{name:hashlib.sha256(Path(result[name]).read_bytes()).hexdigest() for name in ('xlsx','csv','json','coverage')}}
    finally:workbook.close()
    (ROOT/'logs/continuity-export-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'checked_at':report['checked_at'],'jobs':len(expected),'xlsx_csv_json_ids_match':True,'xlsx':result['xlsx']},ensure_ascii=True))


if __name__=='__main__':main()
